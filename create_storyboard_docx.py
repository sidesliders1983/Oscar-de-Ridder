from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

root = Path(__file__).parent
src = root / 'storyboard/book-3/voorleestekst-v001.md'
out = root / 'storyboard/book-3/voorleestekst-v001.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8); sec.right_margin = Inches(0.8)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(11)
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
for name,size in [('Title',24),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
    styles[name].font.name='Aptos Display'; styles[name].font.size=Pt(size); styles[name].font.color.rgb=RGBColor(0,0,0)
    styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos Display')

def add_line(text, style=None):
    p=doc.add_paragraph(style=style)
    p.paragraph_format.space_after=Pt(4)
    if text.startswith('**') and text.endswith('**'):
        r=p.add_run(text[2:-2]); r.bold=True
    else:
        # basic markdown bold
        pos=0
        for m in re.finditer(r'\*\*(.+?)\*\*', text):
            p.add_run(text[pos:m.start()]); rr=p.add_run(m.group(1)); rr.bold=True; pos=m.end()
        p.add_run(text[pos:])
    return p

lines=src.read_text(encoding='utf-8').splitlines()
for line in lines:
    if not line.strip():
        continue
    if line.startswith('# '):
        p=doc.add_paragraph(style='Title'); p.add_run(line[2:].strip()); p.paragraph_format.space_after=Pt(10)
    elif line.startswith('## '):
        add_line(line[3:].strip(),'Heading 1')
    elif line.startswith('### '):
        add_line(line[4:].strip(),'Heading 2')
    elif line.startswith('**') and line.endswith('**'):
        add_line(line,'Heading 3')
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); p.add_run(line[2:])
    else:
        add_line(line.replace('  ',' '))

doc.core_properties.title='Gebulder in de Bergen Voorleestekst'
doc.core_properties.subject='Bewerkbare conceptteksten per storyboardspread'
doc.core_properties.author='Oscar de Ridder'
doc.save(out)
print(out)
